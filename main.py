"""Telegram bot entry point untuk aiagentgres.

Phase 0 security: setiap handler di-guard `require_owner` decorator yang reject
pesan dari user di luar `OWNER_TELEGRAM_IDS`. Sebelumnya cuma EXECUTE_BASH
yang di-check di agent.py, tapi action lain (SKILL, SCREENSHOT, AUTO_*) gak
di-filter, jadi sembarang user bisa trigger LLM untuk minta action sensitif.
"""

import datetime
import os
import tempfile
from functools import wraps

import pytz
from apscheduler.schedulers.asyncio import AsyncIOScheduler
from telegram import Update
from telegram.ext import (
    ApplicationBuilder,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    filters,
)

import auth
import auto_trader
from agent import execute_approved_command, run_agent
from config import TELEGRAM_TOKEN
from logging_setup import configure_logging, get_logger
from memory import clear_history, get_all_memories_text
from morning_briefing import generate_morning_briefing
from paths import TMP_DIR, ensure_runtime_dirs

configure_logging()
ensure_runtime_dirs()
logger = get_logger(__name__)


def require_owner(handler):
    """Decorator yang reject pesan dari user di luar OWNER_TELEGRAM_IDS.

    Wrap setiap Telegram handler. Logging deny attempt buat audit trail.
    """

    @wraps(handler)
    async def wrapper(update: Update, context: ContextTypes.DEFAULT_TYPE):
        user = update.effective_user
        user_id = user.id if user else None
        if not auth.is_authorized(user_id):
            username = getattr(user, "username", None) or "?"
            logger.warning(
                "DENY: user_id=%s username=%s attempt %s",
                user_id,
                username,
                handler.__name__,
            )
            if update.message:
                try:
                    await update.message.reply_text(
                        "⛔ Bot ini private — akses ditolak. "
                        "Hubungi owner kalau lo butuh akses."
                    )
                except Exception:
                    logger.exception("Failed to send deny message")
            elif update.callback_query:
                try:
                    await update.callback_query.answer("⛔ Akses ditolak", show_alert=True)
                except Exception:
                    logger.exception("Failed to answer denied callback")
            return None
        return await handler(update, context)

    return wrapper


@require_owner
async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "Halo! Gua personal AI assistant lo 🤖\n\n"
        "Gua bisa bantu:\n"
        "📧 Kirim & baca email\n"
        "📅 Buat event di Google Calendar\n"
        "🔍 Research & web search\n"
        "🎙️ Voice note\n"
        "📄 Baca PDF, Word, Excel\n"
        "🧠 Inget semua tentang lo\n\n"
        "Commands:\n"
        "/start - Mulai\n"
        "/clear - Hapus history chat\n"
        "/memory - Lihat semua memory\n\n"
        "Ketik, kirim voice note, atau upload file!"
    )


async def send_morning_briefing(app):
    """Kirim morning briefing ke semua owner yang terdaftar."""
    if not auth.OWNER_IDS:
        logger.warning("Skip morning briefing — OWNER_TELEGRAM_IDS belum di-set.")
        return
    try:
        briefing = await generate_morning_briefing()
    except Exception:
        logger.exception("Failed to generate morning briefing")
        return
    for owner_id in auth.OWNER_IDS:
        try:
            await app.bot.send_message(chat_id=owner_id, text=briefing, parse_mode="Markdown")
            logger.info("Morning briefing terkirim ke user_id=%s", owner_id)
        except Exception:
            logger.exception("Gagal kirim morning briefing ke user_id=%s", owner_id)


@require_owner
async def handle_trade_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    data = query.data

    if data.startswith("trade_yes_"):
        parts = data.split("_")
        symbol = parts[2]
        side = parts[3]
        sl = float(parts[4])
        tp = float(parts[5])
        qty = float(parts[6])

        await query.edit_message_text(f"⏳ Executing {side} {symbol}...")

        from tools.binance_trader import open_position

        result = open_position(
            symbol=symbol,
            side=side,
            usdt_amount=qty,
            sl_price=sl,
            tp_price=tp,
            leverage=5,
        )

        if "error" in result:
            await query.edit_message_text(f"❌ Error: {result['error']}")
        else:
            await query.edit_message_text(
                f"✅ <b>ORDER EXECUTED!</b>\n\n"
                f"📊 {result['symbol']} {result['side']}\n"
                f"💰 Entry: ${result['entry']:,}\n"
                f"🛑 SL: ${result['sl']:,}\n"
                f"✅ TP: ${result['tp']:,}\n"
                f"📦 Qty: {result['qty']}",
                parse_mode="HTML",
            )

    elif data.startswith("trade_no_"):
        symbol = data.split("_")[2]
        await query.edit_message_text(f"❌ Trade {symbol} di-skip.")


@require_owner
async def handle_bash_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Callback handler buat tombol Approve/Skip pada pending EXECUTE_BASH."""
    query = update.callback_query
    user_id = query.from_user.id if query and query.from_user else None
    data = query.data if query else ""
    await query.answer()

    if data.startswith("bash_approve:"):
        approval_id = data.split(":", 1)[1]
        await query.edit_message_text("⏳ Approved — eksekusi...")
        reply = await execute_approved_command(approval_id, user_id, context)
        await context.bot.send_message(chat_id=user_id, text=reply)
    elif data.startswith("bash_skip:"):
        approval_id = data.split(":", 1)[1]
        if auth.cancel_pending_command(approval_id, user_id):
            await query.edit_message_text("❌ Command di-skip.")
        else:
            await query.edit_message_text("⚠️ Approval gak ketemu (mungkin udah expired).")
    else:
        logger.warning("Unknown bash callback data: %r", data)


@require_owner
async def clear(update: Update, context: ContextTypes.DEFAULT_TYPE):
    clear_history(update.effective_user.id)
    await update.message.reply_text("✅ History chat udah dihapus!")


@require_owner
async def show_memory(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    await update.message.reply_text(get_all_memories_text(user_id))


@require_owner
async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    user_message = update.message.text

    os.environ["USER_CHAT"] = user_message
    waktu_sekarang = datetime.datetime.now().strftime("%H:%M WIB (Tanggal %d %B %Y)")
    user_message = f"{user_message}\n\n[Sistem: Waktu saat ini {waktu_sekarang}]"

    await update.message.reply_chat_action("typing")
    try:
        reply = await run_agent(user_id, user_message, context)
        await update.message.reply_text(reply)
    except Exception as e:
        logger.exception("handle_message failed for user_id=%s", user_id)
        await update.message.reply_text(f"❌ Terjadi error: {str(e)}")


@require_owner
async def handle_voice(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    await update.message.reply_chat_action("typing")
    try:
        voice = update.message.voice
        file = await context.bot.get_file(voice.file_id)
        with tempfile.NamedTemporaryFile(suffix=".ogg", delete=False) as tmp:
            tmp_path = tmp.name
        await file.download_to_drive(tmp_path)

        from groq import Groq

        from config import GROQ_API_KEY

        groq_client = Groq(api_key=GROQ_API_KEY)
        with open(tmp_path, "rb") as audio_file:
            transcription = groq_client.audio.transcriptions.create(
                file=("voice.ogg", audio_file),
                model="whisper-large-v3",
                language="id",
            )
        os.unlink(tmp_path)

        user_message = transcription.text
        await update.message.reply_text(f'🎙️ Lo bilang: "{user_message}"')
        reply = await run_agent(user_id, user_message, context)
        await update.message.reply_text(reply)
    except Exception as e:
        logger.exception("handle_voice failed for user_id=%s", user_id)
        await update.message.reply_text(f"❌ Gagal proses voice note: {str(e)}")


@require_owner
async def handle_photo(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    await update.message.reply_chat_action("typing")
    try:
        photo = update.message.photo[-1]
        file = await context.bot.get_file(photo.file_id)
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        tmp_path = str(TMP_DIR / f"photo_{user_id}_{timestamp}.jpg")
        await file.download_to_drive(tmp_path)
        caption = update.message.caption or ""
        user_message = caption if caption else "tolong analisa gambar ini"
        reply = await run_agent(user_id, user_message, context)
        await update.message.reply_text(reply)
    except Exception as e:
        logger.exception("handle_photo failed for user_id=%s", user_id)
        await update.message.reply_text(f"❌ Gagal proses gambar: {str(e)}")


@require_owner
async def switch_mode(update: Update, context: ContextTypes.DEFAULT_TYPE):
    args = context.args
    if not args or args[0] not in ["scalping", "swing"]:
        await update.message.reply_text(
            "⚙️ Cara pakai:\n"
            "/mode scalping\n"
            "/mode swing"
        )
        return

    mode = args[0]
    auto_trader.TRADING_MODE = mode

    if mode == "scalping":
        auto_trader.TIMEFRAME = "15m"
        auto_trader.ATR_MULT = 1.5
        auto_trader.RSI_LONG_MIN = 45
        auto_trader.RSI_LONG_MAX = 65
        auto_trader.RSI_SHORT_MIN = 35
        auto_trader.RSI_SHORT_MAX = 55
        auto_trader.CHECK_INTERVAL = 300
        auto_trader.MODE_LABEL = "⚡ SCALPING (15m)"
    else:
        auto_trader.TIMEFRAME = "4h"
        auto_trader.ATR_MULT = 2.5
        auto_trader.RSI_LONG_MIN = 50
        auto_trader.RSI_LONG_MAX = 70
        auto_trader.RSI_SHORT_MIN = 30
        auto_trader.RSI_SHORT_MAX = 50
        auto_trader.CHECK_INTERVAL = 3600
        auto_trader.MODE_LABEL = "🌊 SWING (4h)"

    await update.message.reply_text(
        f"✅ Mode berhasil diganti ke <b>{auto_trader.MODE_LABEL}</b>",
        parse_mode="HTML",
    )


@require_owner
async def handle_file(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    await update.message.reply_chat_action("typing")

    try:
        if update.message.document:
            doc = update.message.document
            file_name = doc.file_name.lower()
            file = await context.bot.get_file(doc.file_id)
        else:
            await update.message.reply_text("❌ Format file tidak didukung.")
            return

        if not any(file_name.endswith(ext) for ext in [".pdf", ".docx", ".xlsx", ".xls"]):
            await update.message.reply_text(
                "❌ File tidak didukung.\n"
                "Format yang bisa dibaca: PDF, Word (.docx), Excel (.xlsx)"
            )
            return

        caption = update.message.caption or ""

        cleanup_keywords = [
            "rapihin",
            "rapiin",
            "bersihkan",
            "clean",
            "cleanup",
            "beresin",
            "perbaiki",
            "format",
            "rapi",
        ]
        is_cleanup = any(kw in caption.lower() for kw in cleanup_keywords)

        if file_name.endswith((".xlsx", ".xls")) and is_cleanup:
            await update.message.reply_text(
                f"📂 File diterima: {doc.file_name}\n⏳ Lagi rapihin Excel lo..."
            )

            with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
                tmp_path = tmp.name
            await file.download_to_drive(tmp_path)

            from tools.file_creator import clean_excel

            output_path = clean_excel(tmp_path, caption)
            os.unlink(tmp_path)

            if isinstance(output_path, str) and output_path.startswith("❌"):
                await update.message.reply_text(output_path)
            else:
                with open(output_path, "rb") as f:
                    await context.bot.send_document(
                        chat_id=user_id,
                        document=f,
                        filename=f"rapi_{doc.file_name}",
                        caption="✅ Excel udah dirapiin bro!",
                    )
            return

        await update.message.reply_text(
            f"📂 File diterima: {doc.file_name}\n⏳ Lagi dibaca..."
        )

        with tempfile.NamedTemporaryFile(
            suffix=os.path.splitext(file_name)[1], delete=False
        ) as tmp:
            tmp_path = tmp.name
        await file.download_to_drive(tmp_path)

        extracted_text = ""

        if file_name.endswith(".pdf"):
            import fitz

            doc_pdf = fitz.open(tmp_path)
            for page in doc_pdf:
                extracted_text += page.get_text()
            doc_pdf.close()

        elif file_name.endswith(".docx"):
            from docx import Document

            doc_word = Document(tmp_path)
            for para in doc_word.paragraphs:
                if para.text.strip():
                    extracted_text += para.text + "\n"

        elif file_name.endswith((".xlsx", ".xls")):
            import openpyxl

            wb = openpyxl.load_workbook(tmp_path, data_only=True)
            for sheet in wb.worksheets:
                extracted_text += f"\n[Sheet: {sheet.title}]\n"
                for row in sheet.iter_rows(values_only=True):
                    row_data = [str(cell) if cell is not None else "" for cell in row]
                    if any(cell.strip() for cell in row_data):
                        extracted_text += " | ".join(row_data) + "\n"

        os.unlink(tmp_path)

        if not extracted_text.strip():
            await update.message.reply_text("❌ File kosong atau tidak bisa dibaca.")
            return

        extracted_text = extracted_text[:4000]
        user_message = (
            f"{caption}\n\n[ISI FILE: {doc.file_name}]\n{extracted_text}"
            if caption
            else f"[ISI FILE: {doc.file_name}]\n{extracted_text}"
        )

        reply = await run_agent(user_id, user_message, context)
        await update.message.reply_text(reply)

    except Exception as e:
        logger.exception("handle_file failed for user_id=%s", user_id)
        await update.message.reply_text(f"❌ Gagal baca file: {str(e)}")


@require_owner
async def manual_briefing(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("⏳ Generating morning briefing... tunggu sebentar!")
    briefing = await generate_morning_briefing()
    await update.message.reply_text(briefing, parse_mode="Markdown")


def main():
    if not TELEGRAM_TOKEN:
        raise RuntimeError(
            "TELEGRAM_BOT_TOKEN gak di-set. Copy .env.example ke .env dan isi token-nya."
        )
    if not auth.OWNER_IDS:
        logger.warning(
            "OWNER_TELEGRAM_IDS kosong — bot bakal reject SEMUA pesan. "
            "Set OWNER_TELEGRAM_IDS=<your_telegram_id> di .env."
        )

    from telegram.request import HTTPXRequest

    request = HTTPXRequest(
        connect_timeout=60.0,
        read_timeout=60.0,
        write_timeout=60.0,
        pool_timeout=60.0,
    )
    app = ApplicationBuilder().token(TELEGRAM_TOKEN).request(request).build()

    scheduler = AsyncIOScheduler(timezone=pytz.timezone("Asia/Jakarta"))
    scheduler.add_job(
        send_morning_briefing,
        "cron",
        hour=7,
        minute=0,
        args=[app],
    )
    scheduler.start()
    logger.info("Scheduler morning briefing aktif — jam 07:00 WIB")

    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("clear", clear))
    app.add_handler(CommandHandler("memory", show_memory))
    app.add_handler(CommandHandler("briefing", manual_briefing))
    app.add_handler(CommandHandler("mode", switch_mode))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_message))
    app.add_handler(MessageHandler(filters.VOICE, handle_voice))
    app.add_handler(MessageHandler(filters.PHOTO, handle_photo))
    app.add_handler(MessageHandler(filters.Document.ALL, handle_file))
    app.add_handler(CallbackQueryHandler(handle_trade_callback, pattern="^trade_"))
    app.add_handler(CallbackQueryHandler(handle_bash_callback, pattern="^bash_"))

    logger.info("Bot berjalan...")
    app.run_polling()


if __name__ == "__main__":
    main()
